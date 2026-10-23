"""Relatórios com exportação Excel / CSV / PDF (seção 12.2)."""
import csv
import io

from flask import Blueprint, abort, current_app, render_template, request, send_file
from flask_login import login_required

from ...permissoes import requer_permissao
from ...services import estoque as est
from ...services import pdf as svc_pdf
from ...services.precos import custo_base
from ...utils import (dias_atras_iso, formatar_data, formatar_data_hora, formatar_moeda, formatar_numero, hoje_iso,
                      mascarar_cpf, normalizar)

bp = Blueprint("relatorios", __name__, url_prefix="/relatorios")

RELATORIOS = {
    "vendas": ("Vendas", "Vendas por período, operador e cliente", True),
    "vendas_itens": ("Vendas por item", "Quantidade, valor e margem por item (produtos × serviços)", True),
    "pagamentos": ("Pagamentos", "Conferência com o serviço externo (forma, NSU)", True),
    "caixas": ("Caixas", "Sessões de caixa e diferenças", True),
    "devolucoes": ("Devoluções e trocas", "Por motivo, cliente e operador", True),
    "creditos": ("Créditos de clientes", "Gerados, utilizados e saldo em aberto", True),
    "estoque": ("Posição de estoque", "Saldo, custo e valor por item", False),
    "estoque_local": ("Estoque por localização", "Saldo por endereço e lote", False),
    "movimentacoes": ("Movimentações", "Compra, venda, devolução, troca, sucata, doação, ajuste", True),
    "perdas": ("Perdas e doações", "Sucata, ajustes negativos e doações ao custo", True),
    "sugestao": ("Sugestão de compra", "Itens em alerta por fornecedor principal", False),
    "validade": ("Validade", "Lotes vencidos e a vencer", False),
    "precos": ("Histórico de preços", "Variações de compra e venda", True),
    "abc": ("Curva ABC", "Classificação por valor vendido", True),
}

# formatos de coluna: t=texto, m=moeda(centavos), q=quantidade, d=data, dh=data/hora, p=percentual, n=inteiro
FMT = {"m": formatar_moeda, "q": lambda v: formatar_numero(v), "d": formatar_data, "dh": formatar_data_hora,
       "p": lambda v: "" if v is None else f"{formatar_numero(v, 2)}%", "n": lambda v: "" if v is None else str(v),
       "t": lambda v: "" if v is None else str(v)}


def _periodo():
    return request.args.get("inicio") or dias_atras_iso(30), request.args.get("fim") or hoje_iso()


def _nomes(repos):
    return ({u["id"]: u["nome"] for u in repos.usuarios.listar()},
            {c["id"]: c for c in repos.clientes.listar()},
            {m["id"]: m for m in repos.materiais.listar()},
            {t["id"]: t["nome"] for t in repos.tipos.listar()},
            {l["id"]: l for l in repos.localizacoes.listar()})


def gerar(chave, repos):
    """Retorna {colunas: [(titulo, fmt)], linhas: [[...]], totais: [...]|None, filtros: [...]}"""
    usuarios, clientes, materiais, tipos, locs = _nomes(repos)
    inicio, fim = _periodo()
    args = request.args
    no_periodo = lambda d: inicio <= d[:10] <= fim  # noqa: E731

    if chave == "vendas":
        status = args.get("status", "")
        vendas = sorted(repos.vendas.listar(lambda v: no_periodo(v["data_hora"]) and (not status or v["status"] == status)),
                        key=lambda v: v["data_hora"])
        linhas = [[v["data_hora"], v["id"], v["status"], usuarios.get(v["usuario_id"], ""),
                   clientes.get(v.get("cliente_id"), {}).get("nome", ""), len(v["itens"]), v["subtotal"],
                   v["desconto_total"], v["total"], ", ".join(svc_pdf.FORMAS.get(p["forma"], p["forma"]) for p in v["pagamentos"])]
                  for v in vendas]
        fin = [v for v in vendas if v["status"] == "finalizada"]
        tot = sum(v["total"] for v in fin)
        return {"colunas": [("Data/hora", "dh"), ("Venda", "t"), ("Situação", "t"), ("Operador", "t"), ("Cliente", "t"),
                            ("Itens", "n"), ("Subtotal", "m"), ("Desconto", "m"), ("Total", "m"), ("Pagamento", "t")],
                "linhas": linhas, "filtros": ["periodo", "status"],
                "totais": ["", f"{len(fin)} finalizadas", "", "", "", "", "", "", tot,
                           f"Ticket médio {formatar_moeda(tot // len(fin)) if fin else '—'}"]}

    if chave == "vendas_itens":
        natureza = args.get("natureza", "")
        agg = {}
        for v in repos.vendas.listar(lambda v: no_periodo(v["data_hora"]) and v["status"] == "finalizada"):
            for it in v["itens"]:
                if natureza and it["natureza"] != natureza:
                    continue
                a = agg.setdefault(it["material_id"], {"qtd": 0, "valor": 0, "dev": 0})
                a["qtd"] += it["quantidade_base"]
                a["valor"] += it["valor_liquido"]
                a["dev"] += it.get("devolvido", 0)
        linhas = []
        for mid, a in sorted(agg.items(), key=lambda x: -x[1]["valor"]):
            m = materiais.get(mid, {})
            custo = (m.get("custo_medio") or custo_base(m) or 0) if m.get("natureza") == "produto" else 0
            custo_total = round(a["qtd"] * custo)
            margem = a["valor"] - custo_total if custo else None
            linhas.append([mid, m.get("nome", ""), tipos.get(m.get("tipo_id"), ""),
                           "Serviço" if m.get("natureza") == "servico" else "Produto", a["qtd"], a["dev"], a["valor"],
                           custo_total if custo else None, margem,
                           round(margem * 100 / a["valor"], 1) if margem is not None and a["valor"] else None])
        return {"colunas": [("Código", "t"), ("Item", "t"), ("Tipo", "t"), ("Natureza", "t"), ("Qtd vendida", "q"),
                            ("Qtd devolvida", "q"), ("Valor líquido", "m"), ("Custo (médio atual)", "m"), ("Margem", "m"),
                            ("Margem %", "p")],
                "linhas": linhas, "filtros": ["periodo", "natureza"],
                "totais": ["", "", "", "", "", "", sum(l[6] for l in linhas), sum(l[7] or 0 for l in linhas),
                           sum(l[8] or 0 for l in linhas), ""]}

    if chave == "pagamentos":
        forma = args.get("forma", "")
        linhas = []
        docs = [(v, "Venda") for v in repos.vendas.listar(lambda v: no_periodo(v["data_hora"]) and v["status"] == "finalizada")]
        docs += [(d, "Troca") for d in repos.devolucoes.listar(lambda d: no_periodo(d["data_hora"]))]
        por_forma = {}
        for doc, origem in sorted(docs, key=lambda x: x[0]["data_hora"]):
            for p in doc.get("pagamentos", []):
                if forma and p["forma"] != forma:
                    continue
                linhas.append([doc["data_hora"], f"{origem} {doc['id']}", svc_pdf.FORMAS.get(p["forma"], p["forma"]),
                               p.get("parcelas"), p.get("nsu", ""), p["valor"]])
                por_forma[p["forma"]] = por_forma.get(p["forma"], 0) + p["valor"]
        resumo = " · ".join(f"{svc_pdf.FORMAS.get(k, k)}: {formatar_moeda(v)}" for k, v in por_forma.items())
        return {"colunas": [("Data/hora", "dh"), ("Documento", "t"), ("Forma", "t"), ("Parcelas", "n"), ("NSU/Autorização", "t"),
                            ("Valor recebido", "m")], "linhas": linhas, "filtros": ["periodo", "forma"],
                "totais": ["", resumo, "", "", "", sum(l[5] for l in linhas)],
                "nota": "Dinheiro: valor recebido antes do troco."}

    if chave == "caixas":
        from ...services.vendas import resumo_caixa
        linhas = []
        for c in sorted(repos.caixas.listar(lambda c: no_periodo(c["abertura"])), key=lambda c: c["abertura"]):
            r = resumo_caixa(repos, c)
            linhas.append([c["id"], usuarios.get(c["usuario_id"], ""), c["abertura"], c.get("fechamento"), c["status"],
                           r["qtd_vendas"], r["total_vendas"], r["esperado"], c.get("valor_contado"), r["diferenca"]])
        return {"colunas": [("Caixa", "t"), ("Operador", "t"), ("Abertura", "dh"), ("Fechamento", "dh"), ("Situação", "t"),
                            ("Vendas", "n"), ("Total vendido", "m"), ("Dinheiro esperado", "m"), ("Contado", "m"), ("Diferença", "m")],
                "linhas": linhas, "filtros": ["periodo"], "totais": None}

    if chave == "devolucoes":
        linhas = [[d["data_hora"], d["id"], "Troca" if d["tipo"] == "troca" else "Devolução", d["venda_id"],
                   clientes.get(d["cliente_id"], {}).get("nome", ""), d["motivo"], usuarios.get(d["usuario_id"], ""),
                   d["valor_devolvido"], d["valor_novos"], d["diferenca"], d["credito_gerado"]]
                  for d in sorted(repos.devolucoes.listar(lambda d: no_periodo(d["data_hora"])), key=lambda d: d["data_hora"])]
        return {"colunas": [("Data/hora", "dh"), ("Documento", "t"), ("Tipo", "t"), ("Venda", "t"), ("Cliente", "t"),
                            ("Motivo", "t"), ("Operador", "t"), ("Devolvido", "m"), ("Itens novos", "m"), ("Diferença", "m"),
                            ("Crédito gerado", "m")], "linhas": linhas, "filtros": ["periodo"],
                "totais": ["", f"{len(linhas)} docs", "", "", "", "", "", sum(l[7] for l in linhas), sum(l[8] for l in linhas),
                           "", sum(l[10] for l in linhas)]}

    if chave == "creditos":
        agg = {}
        for c in repos.creditos.listar():
            a = agg.setdefault(c["cliente_id"], {"ger": 0, "uso": 0})
            if no_periodo(c["data_hora"]):
                if c["tipo"] == "uso":
                    a["uso"] += c["valor"]
                elif c["valor"] > 0:
                    a["ger"] += c["valor"]
                else:
                    a["uso"] += -c["valor"]
        saldos = repos.creditos.saldos()
        linhas = [[cid, clientes.get(cid, {}).get("nome", ""), mascarar_cpf(clientes.get(cid, {}).get("cpf")),
                   a["ger"], a["uso"], saldos.get(cid, 0)]
                  for cid, a in sorted(agg.items(), key=lambda x: -saldos.get(x[0], 0))
                  if a["ger"] or a["uso"] or saldos.get(cid)]
        return {"colunas": [("Cliente", "t"), ("Nome", "t"), ("CPF", "t"), ("Gerado no período", "m"),
                            ("Utilizado no período", "m"), ("Saldo atual", "m")], "linhas": linhas, "filtros": ["periodo"],
                "totais": ["", "", "", sum(l[3] for l in linhas), sum(l[4] for l in linhas), sum(l[5] for l in linhas)]}

    if chave == "estoque":
        tipo, fornecedor, status = args.get("tipo", ""), args.get("fornecedor", ""), args.get("status_estoque", "")
        linhas = []
        for p in sorted(est.posicao(repos), key=lambda p: normalizar(p["material"]["nome"])):
            m = p["material"]
            if tipo and m.get("tipo_id") != tipo:
                continue
            if fornecedor and not any(l["fornecedor_id"] == fornecedor for l in m.get("fornecedores", [])):
                continue
            if status and p["status"] != status:
                continue
            linhas.append([m["id"], m["nome"], tipos.get(m.get("tipo_id"), ""), "m" if m.get("unidade") == "metro" else "un",
                           p["saldo"], m.get("estoque_minimo"), {"ok": "OK", "baixo": "Baixo", "zerado": "Zerado"}[p["status"]],
                           p["custo_unitario"], p["valor_custo"], m.get("preco_venda"), p["valor_venda"]])
        return {"colunas": [("Código", "t"), ("Item", "t"), ("Tipo", "t"), ("Un.", "t"), ("Saldo", "q"), ("Mínimo", "q"),
                            ("Situação", "t"), ("Custo médio", "m"), ("Valor custo", "m"), ("Preço venda", "m"),
                            ("Valor venda", "m")], "linhas": linhas, "filtros": ["tipo", "fornecedor", "status_estoque"],
                "totais": ["", f"{len(linhas)} itens", "", "", "", "", "", "", sum(l[8] for l in linhas), "",
                           sum(l[10] for l in linhas)]}

    if chave == "estoque_local":
        loc_id = args.get("loc", "")
        lotes = {l["id"]: l for l in repos.lotes.listar()}
        linhas = []
        for s in repos.saldos.listar(lambda s: not loc_id or s["localizacao_id"] == loc_id):
            m = materiais.get(s["material_id"], {})
            lote = lotes.get(s.get("lote_id"))
            linhas.append([locs.get(s["localizacao_id"], {}).get("codigo", ""), m.get("id"), m.get("nome", ""),
                           lote.get("validade") if lote else None, s["quantidade"],
                           round(s["quantidade"] * (m.get("custo_medio") or 0))])
        linhas.sort(key=lambda l: (l[0], normalizar(l[2])))
        return {"colunas": [("Localização", "t"), ("Código", "t"), ("Item", "t"), ("Validade do lote", "d"), ("Saldo", "q"),
                            ("Valor custo", "m")], "linhas": linhas, "filtros": ["loc"],
                "totais": ["", "", "", "", "", sum(l[5] for l in linhas)]}

    if chave in ("movimentacoes", "perdas"):
        tipo = args.get("tipo_mov", "")
        if chave == "perdas":
            filtro = lambda m: m["sentido"] == "saida" and m["tipo"] in ("sucata", "doacao", "ajuste")  # noqa: E731
        else:
            filtro = lambda m: not tipo or m["tipo"] == tipo  # noqa: E731
        lotes = {l["id"]: l for l in repos.lotes.listar()}
        linhas = []
        for m in sorted(repos.movimentacoes.listar(inicio, fim, filtro), key=lambda m: (m["data_hora"], m["id"])):
            mat = materiais.get(m["material_id"], {})
            lote = lotes.get(m.get("lote_id"))
            ref = m.get("venda_id") or m.get("devolucao_id") or m.get("compra_id") or m.get("inventario_id") or m.get("transferencia_id") or ""
            obs = " ".join(x for x in (m.get("motivo"), m.get("destinatario"), m.get("documento"), m.get("observacao")) if x)
            if chave == "perdas":
                custo = m.get("custo_unitario") or 0
                linhas.append([m["data_hora"], est.TIPOS_MOV.get(m["tipo"], m["tipo"]), mat.get("id"), mat.get("nome", ""),
                               m["quantidade"], custo, round(m["quantidade"] * custo), obs])
            else:
                linhas.append([m["data_hora"], m["id"], est.TIPOS_MOV.get(m["tipo"], m["tipo"]),
                               "Entrada" if m["sentido"] == "entrada" else "Saída", mat.get("id"), mat.get("nome", ""),
                               locs.get(m["localizacao_id"], {}).get("codigo", ""), lote.get("validade") if lote else None,
                               m["quantidade"] if m["sentido"] == "entrada" else -m["quantidade"], m.get("valor_unitario"),
                               ref, usuarios.get(m.get("usuario_id"), ""), obs])
        if chave == "perdas":
            return {"colunas": [("Data/hora", "dh"), ("Tipo", "t"), ("Código", "t"), ("Item", "t"), ("Qtd", "q"),
                                ("Custo unit.", "m"), ("Valor ao custo", "m"), ("Motivo / destinatário", "t")],
                    "linhas": linhas, "filtros": ["periodo"],
                    "totais": ["", "", "", "", "", "", sum(l[6] for l in linhas), ""]}
        return {"colunas": [("Data/hora", "dh"), ("Mov.", "t"), ("Tipo", "t"), ("Sentido", "t"), ("Código", "t"), ("Item", "t"),
                            ("Local", "t"), ("Validade", "d"), ("Qtd", "q"), ("Valor unit.", "m"), ("Referência", "t"),
                            ("Usuário", "t"), ("Obs.", "t")], "linhas": linhas, "filtros": ["periodo", "tipo_mov"], "totais": None}

    if chave == "sugestao":
        linhas = []
        for g in est.sugestao_compra(repos):
            f = g["fornecedor"]
            for it in g["itens"]:
                m = it["material"]
                linhas.append([f.get("fantasia") or f["razao_social"], m["id"], m["nome"], it["saldo"], m.get("estoque_minimo"),
                               it["sugerido"], it["caixas"], it["linha"].get("preco_compra"),
                               "caixa" if it["linha"].get("unidade_compra") == "caixa" else "item"])
        return {"colunas": [("Fornecedor", "t"), ("Código", "t"), ("Item", "t"), ("Saldo", "q"), ("Mínimo", "q"),
                            ("Sugerido (un/m)", "q"), ("Caixas", "n"), ("Último preço", "m"), ("Por", "t")],
                "linhas": linhas, "filtros": [], "totais": None}

    if chave == "validade":
        dados = est.alertas_validade(repos, int(args.get("dias") or 30))
        linhas = [[st, i["material"].get("id"), i["material"].get("nome", ""), i["lote"]["validade"],
                   i["localizacao"].get("codigo", ""), i["saldo"]["quantidade"]]
                  for st, lst in (("Vencido", dados["vencidos"]), ("A vencer", dados["a_vencer"])) for i in lst]
        return {"colunas": [("Situação", "t"), ("Código", "t"), ("Item", "t"), ("Validade", "d"), ("Local", "t"), ("Saldo", "q")],
                "linhas": linhas, "filtros": ["dias"], "totais": None}

    if chave == "precos":
        forn = {f["id"]: f.get("fantasia") or f["razao_social"] for f in repos.fornecedores.listar()}
        campos = {"preco_venda": "Venda", "preco_venda_caixa": "Venda (caixa)", "preco_compra": "Compra"}
        linhas = [[h["data_hora"], h["material_id"], materiais.get(h["material_id"], {}).get("nome", ""),
                   campos.get(h["campo"], h["campo"]), forn.get(h.get("fornecedor_id"), ""), h["anterior"], h["novo"],
                   h["var_abs"], h["var_pct"], h.get("origem", ""), usuarios.get(h.get("usuario_id"), "")]
                  for h in sorted(repos.historico_precos.listar(lambda h: no_periodo(h["data_hora"])), key=lambda h: h["data_hora"])]
        return {"colunas": [("Data/hora", "dh"), ("Código", "t"), ("Item", "t"), ("Preço", "t"), ("Fornecedor", "t"),
                            ("Anterior", "m"), ("Novo", "m"), ("Variação R$", "m"), ("Variação %", "p"), ("Origem", "t"),
                            ("Usuário", "t")], "linhas": linhas, "filtros": ["periodo"], "totais": None}

    if chave == "abc":
        agg = {}
        for v in repos.vendas.listar(lambda v: no_periodo(v["data_hora"]) and v["status"] == "finalizada"):
            for it in v["itens"]:
                agg[it["material_id"]] = agg.get(it["material_id"], 0) + it["valor_liquido"]
        total = sum(agg.values()) or 1
        linhas, acum = [], 0
        for mid, valor in sorted(agg.items(), key=lambda x: -x[1]):
            acum += valor
            pct_acum = acum * 100 / total
            classe = "A" if pct_acum - valor * 100 / total < 80 else "B" if pct_acum - valor * 100 / total < 95 else "C"
            linhas.append([classe, mid, materiais.get(mid, {}).get("nome", ""), valor, round(valor * 100 / total, 2),
                           round(pct_acum, 2)])
        return {"colunas": [("Classe", "t"), ("Código", "t"), ("Item", "t"), ("Valor vendido", "m"), ("% do total", "p"),
                            ("% acumulado", "p")], "linhas": linhas, "filtros": ["periodo"],
                "totais": ["", "", "", sum(agg.values()), "", ""]}
    abort(404)


def _fmt(f, v):
    if v is None or v == "":
        return ""
    if isinstance(v, str) and f not in ("dh", "d"):
        return v
    return FMT[f](v)


def _formatar(rel):
    fmts = [f for _, f in rel["colunas"]]
    linhas = [[_fmt(f, v) for f, v in zip(fmts, l)] for l in rel["linhas"]]
    totais = [_fmt(f, v) for f, v in zip(fmts, rel["totais"])] if rel.get("totais") else None
    return linhas, totais


@bp.route("/")
@login_required
@requer_permissao("relatorios.visualizar")
def index():
    return render_template("relatorios/index.html", relatorios=RELATORIOS)


@bp.route("/<chave>")
@login_required
@requer_permissao("relatorios.visualizar")
def ver(chave):
    if chave not in RELATORIOS:
        abort(404)
    repos = current_app.repos
    rel = gerar(chave, repos)
    titulo = RELATORIOS[chave][0]
    linhas, totais = _formatar(rel)
    inicio, fim = _periodo()
    subtitulo = f"Período: {formatar_data(inicio)} a {formatar_data(fim)}" if RELATORIOS[chave][2] else f"Posição em {formatar_data(hoje_iso())}"
    nome = f"{chave}_{hoje_iso()}"
    formato = request.args.get("formato", "")
    colunas = [c for c, _ in rel["colunas"]]
    if formato == "csv":
        buf = io.StringIO()
        w = csv.writer(buf, delimiter=";")
        w.writerow(colunas)
        w.writerows(linhas)
        if totais:
            w.writerow(totais)
        return send_file(io.BytesIO(buf.getvalue().encode("utf-8-sig")), mimetype="text/csv", as_attachment=True,
                         download_name=f"{nome}.csv")
    if formato == "xlsx":
        return send_file(_xlsx(titulo, subtitulo, rel), as_attachment=True, download_name=f"{nome}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    if formato == "pdf":
        buf = io.BytesIO()
        svc_pdf.relatorio_tabela(buf, titulo, colunas, linhas, subtitulo, totais)
        buf.seek(0)
        return send_file(buf, mimetype="application/pdf", as_attachment=True, download_name=f"{nome}.pdf")
    return render_template("relatorios/ver.html", chave=chave, titulo=titulo, descricao=RELATORIOS[chave][1],
                           subtitulo=subtitulo, colunas=rel["colunas"], linhas=linhas, totais=totais, nota=rel.get("nota"),
                           filtros=rel["filtros"], args=request.args, inicio=inicio, fim=fim,
                           tipos=repos.tipos.listar(), fornecedores=repos.fornecedores.listar(),
                           localizacoes=sorted(repos.localizacoes.listar(), key=lambda l: l["codigo"]))


def _xlsx(titulo, subtitulo, rel):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:31]
    ws.append([titulo])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([subtitulo])
    ws.append([])
    ws.append([c for c, _ in rel["colunas"]])
    for cel in ws[4]:
        cel.font = Font(bold=True)
        cel.fill = PatternFill("solid", fgColor="E9ECEF")
    fmts = [f for _, f in rel["colunas"]]

    def conv(f, v):
        if v is None or v == "":
            return None
        if f == "m" and isinstance(v, (int, float)):
            return v / 100
        if f == "dh" and isinstance(v, str):
            return formatar_data_hora(v)
        if f == "d" and isinstance(v, str):
            return formatar_data(v)
        return v
    for l in rel["linhas"] + ([rel["totais"]] if rel.get("totais") else []):
        ws.append([conv(f, v) for f, v in zip(fmts, l)])
        for i, f in enumerate(fmts, 1):
            c = ws.cell(row=ws.max_row, column=i)
            if f == "m":
                c.number_format = '"R$" #,##0.00'
            elif f == "q":
                c.number_format = "#,##0.###"
    if rel.get("totais"):
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
    for i, (c, _) in enumerate(rel["colunas"], 1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(45, len(c) + 4))
    ws.freeze_panes = "A5"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
